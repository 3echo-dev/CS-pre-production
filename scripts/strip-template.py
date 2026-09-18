#!/usr/bin/env python3
"""Strip a client's finished workbook down to an empty template, and say what was kept.

    python strip-template.py --scan <file.xlsx> [--json]
    python strip-template.py <source.xlsx> --out <template.xlsx> [--sheet NAME] (--header N | --layout)
                             [--clear A5:G28 ...] [--force] [--json]
    python strip-template.py --accept <template.xlsx> --by "<who>"

The realistic way a client supplies a template is by handing over a finished job's workbook,
and a filled workbook is not a template: it hands the compiler the answers it is supposed to
produce, and a filled call sheet carries crew phone numbers into every job folder that uses it.

--scan reads every sheet, finds the header row (the row with the most text cells among the
first twenty), counts the rows below it that hold anything, the images, the sheets and the
size, and says whether the file is clean, filled or accepted. A sidecar <file>.source.json,
written by a strip or an --accept, whose hash still matches the file, is what makes it
accepted; the heuristic alone is a warning, because the header block is where the leaks hide.

A strip keeps one sheet and drops the others, drops every image and chart, deletes every row
below the header (--header N) or keeps the form as it is (--layout, for a call sheet whose
labels sit in columns), blanks any --clear ranges, then reads the result back and prints every
cell it kept, so a person reads the header block before trusting it. --accept records that a
person read the kept cells and confirmed the file is empty.

Exit 0 clean, stripped or accepted · 1 the file holds data, or --out exists · 2 usage · 3 the file is missing or unreadable
"""
import argparse, datetime, hashlib, json, os, re, sys

try:
    import openpyxl
    from openpyxl.cell.cell import MergedCell
    from openpyxl.utils import range_boundaries
except ImportError:
    print("openpyxl is not installed. Run: python -m pip install openpyxl", file=sys.stderr)
    sys.exit(3)

BIG_BYTES = 5 * 1024 * 1024
PHONE = re.compile(r"(?<!\d)(\+?\d[\d \-]{7,}\d)(?!\d)")

ap = argparse.ArgumentParser(add_help=True)
ap.add_argument("source", nargs="?")
ap.add_argument("--scan")
ap.add_argument("--accept")
ap.add_argument("--by")
ap.add_argument("--out")
ap.add_argument("--sheet")
ap.add_argument("--header", type=int)
ap.add_argument("--layout", action="store_true")
ap.add_argument("--clear", action="append", default=[])
ap.add_argument("--force", action="store_true")
ap.add_argument("--json", action="store_true")
a = ap.parse_args()

USAGE = ("usage: strip-template.py --scan <file.xlsx> [--json] | <source.xlsx> --out <template.xlsx> [--sheet NAME]"
         " (--header N | --layout) [--clear A5:G28 ...] [--force] | --accept <template.xlsx> --by \"<who>\"")

def usage(msg):
    print("Nothing was written. " + msg, file=sys.stderr)
    print(USAGE, file=sys.stderr)
    sys.exit(2)

def sha256(p):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()

def now():
    return datetime.datetime.now().astimezone().isoformat(timespec="minutes")

def sidecar_path(p):
    return p + ".source.json"

def read_sidecar(p):
    try:
        with open(sidecar_path(p), encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return None

def write_sidecar(p, side):
    with open(sidecar_path(p), "w", encoding="utf-8") as f:
        json.dump(side, f, indent=2)
        f.write("\n")

def text(v):
    return "" if v is None else str(v).strip()

def load(p):
    if not os.path.isfile(p):
        print("No such file: " + p, file=sys.stderr)
        sys.exit(3)
    try:
        return openpyxl.load_workbook(p)
    except Exception as e:
        print("Cannot read " + p + ": " + str(e), file=sys.stderr)
        sys.exit(3)

def header_row(ws):
    best, score = 1, -1
    for r in range(1, min(ws.max_row, 20) + 1):
        n = 0
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip():
                n += 1
        if n > score:
            best, score = r, n
    return best

def scan_sheet(ws):
    h = header_row(ws)
    non_empty, phones, data_rows = 0, 0, []
    for r in range(1, ws.max_row + 1):
        raw = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        filled = [text(v) for v in raw if text(v)]
        non_empty += len(filled)
        # A date cell prints as digits and separators, which is what a phone number looks like.
        phones += sum(1 for v in raw if text(v) and not isinstance(v, (datetime.date, datetime.datetime)) and PHONE.search(text(v)))
        if r > h and filled:
            data_rows.append((r, filled))
    return {
        "name": ws.title, "rows": ws.max_row, "columns": ws.max_column, "header": h,
        "nonEmpty": non_empty, "dataRows": len(data_rows),
        "sample": [{"row": r, "cells": f[:6]} for r, f in data_rows[:3]],
        "phoneLike": phones,
        "images": len(getattr(ws, "_images", []) or []),
        "charts": len(getattr(ws, "_charts", []) or []),
        "merged": len(ws.merged_cells.ranges),
    }

def scan(p):
    wb = load(p)
    size = os.path.getsize(p)
    sheets = [scan_sheet(ws) for ws in wb.worksheets]
    side = read_sidecar(p)
    digest = sha256(p)
    accepted = bool(side and side.get("sha256") == digest)
    filled = [s for s in sheets if s["dataRows"]]
    verdict = "accepted" if accepted else ("filled" if filled else "clean")
    warnings = []
    if size > BIG_BYTES:
        warnings.append("%.1f MB; embedded images are the usual cause, and a strip drops them" % (size / 1048576.0))
    if len(sheets) > 1:
        warnings.append("%d sheets; a template is one sheet" % len(sheets))
    imgs = sum(s["images"] for s in sheets)
    if imgs:
        warnings.append("%d embedded image%s" % (imgs, "" if imgs == 1 else "s"))
    ph = sum(s["phoneLike"] for s in sheets)
    if ph and not accepted:
        warnings.append("%d cell%s that read%s like phone numbers" % (ph, "" if ph == 1 else "s", "s" if ph == 1 else ""))
    if side and not accepted:
        warnings.append("the file changed after it was stripped or accepted; check it again")
    return {"file": p.replace(os.sep, "/"), "bytes": size, "sheets": sheets,
            "sidecar": side if accepted else None, "verdict": verdict, "accepted": accepted,
            "warnings": warnings, "sha256": digest}

def say_scan(res):
    n = len(res["sheets"])
    lines = ["%s: %s, %d bytes, %d sheet%s" % (os.path.basename(res["file"]), res["verdict"], res["bytes"], n, "" if n == 1 else "s")]
    for s in res["sheets"]:
        lines.append("  sheet %s: header row %d, %d row%s below it with content, %d filled cells" %
                     (s["name"], s["header"], s["dataRows"], "" if s["dataRows"] == 1 else "s", s["nonEmpty"]))
        for smp in s["sample"]:
            lines.append("    row %d: %s" % (smp["row"], " | ".join(smp["cells"])))
    for w in res["warnings"]:
        lines.append("  warning: " + w)
    return "\n".join(lines)

def kept_cells(p):
    ws = openpyxl.load_workbook(p).active
    out = []
    for row in ws.iter_rows():
        for cell in row:
            v = text(cell.value)
            if v:
                out.append(cell.coordinate + ": " + v)
    return out

# ---------------------------------------------------------------- --scan
if a.scan:
    res = scan(a.scan)
    print(json.dumps(res) if a.json else say_scan(res))
    sys.exit(0 if res["verdict"] != "filled" else 1)

# ---------------------------------------------------------------- --accept
if a.accept:
    if not a.by:
        usage("--accept needs --by \"<who>\": the person who read the kept cells.")
    res = scan(a.accept)
    side = read_sidecar(a.accept) or {}
    side.update({"acceptedBy": a.by, "acceptedAt": now(), "sha256": res["sha256"],
                 "keptCells": sum(s["nonEmpty"] for s in res["sheets"])})
    write_sidecar(a.accept, side)
    print("Accepted %s as empty, by %s: %d cells kept. Recorded in %s." %
          (os.path.basename(a.accept), a.by, side["keptCells"], os.path.basename(sidecar_path(a.accept))))
    sys.exit(0)

# ---------------------------------------------------------------- strip
if not a.source:
    usage("Give a workbook to strip, --scan a template, or --accept one.")
if not a.out:
    usage("--out names the template to write, under client/templates/.")
if bool(a.header) == bool(a.layout):
    usage("Choose one of --header N (a table: rows below N are deleted) or --layout (a form: cells stay; blank them with --clear).")
if not os.path.isfile(a.source):
    print("No such file: " + a.source, file=sys.stderr)
    sys.exit(3)
if os.path.exists(a.out) and not a.force:
    print("Nothing was written. %s exists; --force overwrites it." % a.out, file=sys.stderr)
    sys.exit(1)
wb = load(a.source)
if a.sheet:
    if a.sheet not in wb.sheetnames:
        print("Nothing was written. No sheet named %s. Sheets: %s" % (a.sheet, ", ".join(wb.sheetnames)), file=sys.stderr)
        sys.exit(1)
    ws = wb[a.sheet]
else:
    ws = wb.active
dropped_sheets = [o.title for o in wb.worksheets if o is not ws]
for other in list(wb.worksheets):
    if other is not ws:
        wb.remove(other)
dropped_images = len(getattr(ws, "_images", []) or []) + len(getattr(ws, "_charts", []) or [])
ws._images = []
ws._charts = []
cleared = []
if a.header:
    if a.header > ws.max_row:
        usage("--header %d is past the last row (%d)." % (a.header, ws.max_row))
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row > a.header:
            ws.unmerge_cells(str(rng))
        elif rng.max_row > a.header:
            ws.unmerge_cells(str(rng))
            ws.merge_cells(start_row=rng.min_row, start_column=rng.min_col, end_row=a.header, end_column=rng.max_col)
    if ws.max_row > a.header:
        ws.delete_rows(a.header + 1, ws.max_row - a.header)
for rng in a.clear:
    try:
        c1, r1, c2, r2 = range_boundaries(rng)
    except Exception:
        usage("--clear takes a range like A5:G28, not " + rng)
    n = 0
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(r, c)
            if isinstance(cell, MergedCell) or cell.value is None:
                continue
            cell.value = None
            n += 1
    cleared.append({"range": rng, "cells": n})
os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
wb.save(a.out)
kept = kept_cells(a.out)
side = {"source": os.path.abspath(a.source).replace(os.sep, "/"), "sourceName": os.path.basename(a.source),
        "sourceBytes": os.path.getsize(a.source),
        "sourceModified": datetime.datetime.fromtimestamp(os.path.getmtime(a.source)).astimezone().isoformat(timespec="minutes"),
        "sheet": ws.title, "headerRow": a.header if a.header else "layout", "cleared": cleared,
        "droppedSheets": dropped_sheets, "droppedImages": dropped_images,
        "strippedAt": now(), "keptCells": len(kept), "sha256": sha256(a.out)}
write_sidecar(a.out, side)
if a.json:
    print(json.dumps({"template": a.out.replace(os.sep, "/"), "sidecar": side, "kept": kept}))
    sys.exit(0)
print("Stripped %s sheet %s to %s (%d bytes). Dropped %d other sheet%s and %d image%s; %s%s." % (
    side["sourceName"], ws.title, a.out.replace(os.sep, "/"), os.path.getsize(a.out),
    len(dropped_sheets), "" if len(dropped_sheets) == 1 else "s", dropped_images, "" if dropped_images == 1 else "s",
    ("deleted every row below row %d" % a.header) if a.header else "kept the form as it is",
    ("; cleared %d cells" % sum(c["cells"] for c in cleared)) if cleared else ""))
print("Kept %d cells. Read them before trusting the template; a name, a number or a place here is a leak:" % len(kept))
for k in kept:
    print("  " + k)
print("Recorded where it came from in " + os.path.basename(sidecar_path(a.out)) + ".")
