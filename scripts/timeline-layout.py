#!/usr/bin/env python3
"""Lay a twelve-month calendar timeline template out for one year.

    python timeline-layout.py <template.xlsx> --year 2027 --out <job>/timeline.xlsx

The house timeline (templates/house/timeline.xlsx) has one sheet per month, named Jan..Dec: the
month's first day as a date in A2 (shown as "January 2026"), weekday names in row 5, and a week
every five rows from row 6, Sunday in column A, with four event rows under each day number. This
writes the month dates and day numbers for the year asked, and clears the old ones. Event rows,
the Notes cell, styles and merges are left as they are. A sixth week goes on row 31 beside the
Notes cell, as the template does (a month never needs more than two days there).

Exit 0 wrote · 1 the template is not a month-per-sheet calendar · 2 usage
"""
import argparse, calendar, datetime, sys
import openpyxl

ap = argparse.ArgumentParser()
ap.add_argument('template')
ap.add_argument('--year', type=int, required=True)
ap.add_argument('--out', required=True)
a = ap.parse_args()

MONTHS = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
FIRST_ROW, STEP, WEEKS = 6, 5, 6
wb = openpyxl.load_workbook(a.template)
found = [m for m in MONTHS if m in wb.sheetnames]
if not found:
    print('the template has no sheets named Jan..Dec, so it is not a month-per-sheet calendar', file=sys.stderr)
    sys.exit(1)
cal = calendar.Calendar(firstweekday=6)  # Sunday first, as the template's column A
for m in found:
    ws = wb[m]
    month = MONTHS.index(m) + 1
    ws['A2'] = datetime.datetime(a.year, month, 1)
    # Clear the old day numbers on the week rows only; everything else stays.
    for w in range(WEEKS):
        r = FIRST_ROW + STEP * w
        for c in range(1, 8):
            v = ws.cell(r, c).value
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                ws.cell(r, c).value = None
    weeks = cal.monthdayscalendar(a.year, month)
    for w, days in enumerate(weeks[:WEEKS]):
        r = FIRST_ROW + STEP * w
        for c, d in enumerate(days, start=1):
            if d:
                ws.cell(r, c).value = d
wb.save(a.out)
print('Timeline laid out for %d: %d month sheet%s -> %s' % (a.year, len(found), '' if len(found) == 1 else 's', a.out))
